"""
generate_findings.py - end-to-end assessment run and single project entrypoint.

Pipeline
--------
    generate_data  ->  control_tests  ->  risk_engine  ->  outputs/

Running `python src/generate_findings.py` regenerates the synthetic environment,
runs all 48 controls, scores the 137 findings, computes the readiness score and
writes every deliverable in outputs/.  A validation block asserts the headline
project metrics (2,000 assets / 48 controls / 137 findings) and fails loudly on
any drift.
"""

from __future__ import annotations

import os
from collections import Counter

import pandas as pd

import config
import control_tests
import generate_data
import risk_engine as re

OUT_DIR = os.path.join(os.path.dirname(__file__), os.pardir, "outputs")
PBI_DIR = os.path.join(OUT_DIR, "powerbi")

SEV_ORDER = ["Critical", "High", "Medium", "Low"]
PHASE_ORDER = {"Stabilize": 0, "Govern": 1, "Scale": 2}
PHASE_LABEL = {
    "Stabilize": "Days 0-30: Stabilize",
    "Govern": "Days 31-60: Govern",
    "Scale": "Days 61-90: Scale",
}
SEV_FILL = {
    "Critical": "C0392B", "High": "E67E22", "Medium": "F1C40F", "Low": "7FB3D5",
}

# ---------------------------------------------------------------------------
# Builders for each output table
# ---------------------------------------------------------------------------

def build_risk_register(findings) -> pd.DataFrame:
    rows = []
    for f in findings:
        rows.append({
            "finding_id": f.finding_id,
            "control_id": f.control_id,
            "control_name": f.control_name,
            "domain": f.domain,
            "finding_description": (
                f"{f.test_summary} {f.affected_assets} affected in {f.business_unit}."),
            "affected_assets": f.affected_assets,
            "likelihood": f.likelihood,
            "impact": f.impact,
            "risk_score": f.risk_score,
            "severity": f.severity,
            "business_impact": f.business_impact,
            "recommendation": f.recommendation,
            "owner": f.owner,
            "remediation_effort": f.remediation_effort,
            "estimated_risk_reduction": f.estimated_risk_reduction,
            "nist_csf": f.nist_csf,
            "nist_ai_rmf": f.nist_ai_rmf,
            "osfi_b13": f.osfi_b13,
            "iso_27001": f.iso_27001,
            "target_remediation_days": f.target_remediation_days,
            "business_unit": f.business_unit,
            "initiative": f.initiative,
            "roadmap_phase": f.roadmap_phase,
            "risk_rationale": f.risk_rationale,
            "affected_asset_ids": f.affected_ids,
        })
    df = pd.DataFrame(rows)
    df["severity"] = pd.Categorical(df["severity"], SEV_ORDER, ordered=True)
    return df.sort_values(["severity", "risk_score", "finding_id"],
                          ascending=[True, False, True]).reset_index(drop=True)


def build_assessment_results(results) -> pd.DataFrame:
    rows = []
    for r in results:
        sev = Counter(f.severity for f in r.findings)
        max_sev = next((s for s in SEV_ORDER if sev.get(s)), "None")
        rows.append({
            "control_id": r.control_id,
            "control_name": r.control_name,
            "domain": r.domain,
            "dataset": r.dataset,
            "assets_in_scope": r.in_scope,
            "assets_in_violation": r.violations,
            "status": r.status,
            "findings": len(r.findings),
            "max_severity": max_sev,
            "critical_findings": sev.get("Critical", 0),
            "high_findings": sev.get("High", 0),
            "medium_findings": sev.get("Medium", 0),
            "low_findings": sev.get("Low", 0),
        })
    cat = {c["control_id"]: c for c in config.CONTROL_CATALOG}
    df = pd.DataFrame(rows)
    for col in ["initiative", "owner", "roadmap_phase",
                "nist_csf", "nist_ai_rmf", "osfi_b13", "iso_27001"]:
        df[col] = df["control_id"].map(lambda cid: cat[cid][col])
    return df


def build_framework_crosswalk(results) -> pd.DataFrame:
    rows = []
    for r in results:
        c = next(cc for cc in config.CONTROLS if cc["id"] == r.control_id)
        rows.append({
            "control_id": r.control_id,
            "control_name": r.control_name,
            "domain": r.domain,
            "status": r.status,
            "findings": len(r.findings),
            "NIST CSF 2.0": c["nist_csf"],
            "NIST AI RMF": c["nist_ai_rmf"],
            "OSFI B-13": c["osfi_b13"],
            "ISO/IEC 27001:2022": c["iso_27001"],
        })
    return pd.DataFrame(rows)


def build_powerbi_tables(findings, results, readiness):
    tables = {}

    # findings (flat)
    reg = build_risk_register(findings)
    reg_flat = reg.copy()
    reg_flat["severity"] = reg_flat["severity"].astype(str)
    tables["findings"] = reg_flat

    # severity summary
    sev = Counter(f.severity for f in findings)
    tables["severity_summary"] = pd.DataFrame([
        {"severity": s, "findings": sev.get(s, 0),
         "pct_of_findings": round(100 * sev.get(s, 0) / len(findings), 1)}
        for s in SEV_ORDER
    ])

    # domain summary (readiness + findings + control status)
    status_by_domain = {d: Counter() for d in config.DOMAINS}
    for r in results:
        status_by_domain[r.domain][r.status] += 1
    drows = []
    for d in config.DOMAINS:
        df_ = [f for f in findings if f.domain == d]
        sc = Counter(f.severity for f in df_)
        st = status_by_domain[d]
        drows.append({
            "domain": d,
            "readiness_score": readiness["by_domain"][d],
            "findings": len(df_),
            "critical": sc.get("Critical", 0),
            "high": sc.get("High", 0),
            "medium": sc.get("Medium", 0),
            "low": sc.get("Low", 0),
            "total_risk_score": sum(f.risk_score for f in df_),
            "avg_risk_score": round(sum(f.risk_score for f in df_) / max(len(df_), 1), 1),
            "controls_not_effective": st.get("Not Effective", 0),
            "controls_partially_effective": st.get("Partially Effective", 0),
            "controls_effective": st.get("Effective", 0),
        })
    tables["domain_summary"] = pd.DataFrame(drows)

    # control status
    tables["control_status"] = build_assessment_results(results)[
        ["control_id", "control_name", "domain", "status", "findings",
         "max_severity", "assets_in_scope", "assets_in_violation"]]

    # top 10 risks
    tr = re.top_risks(findings, 10)
    tables["top10_risks"] = pd.DataFrame([{
        "rank": i + 1,
        "finding_id": f.finding_id,
        "control_id": f.control_id,
        "title": f.control_name,
        "domain": f.domain,
        "business_unit": f.business_unit,
        "affected_assets": f.affected_assets,
        "risk_score": f.risk_score,
        "severity": f.severity,
        "business_impact": f.business_impact,
    } for i, f in enumerate(tr)])

    # framework coverage
    fw_cols = {"NIST CSF 2.0": "nist_csf", "NIST AI RMF": "nist_ai_rmf",
               "OSFI B-13": "osfi_b13", "ISO/IEC 27001:2022": "iso_27001"}
    frows = []
    for label, attr in fw_cols.items():
        mapped_controls = [c for c in config.CONTROLS if c[attr] != "-"]
        mapped_ids = {c["id"] for c in mapped_controls}
        frows.append({
            "framework": label,
            "controls_mapped": len(mapped_controls),
            "controls_total": len(config.CONTROLS),
            "coverage_pct": round(100 * len(mapped_controls) / len(config.CONTROLS), 1),
            "findings_mapped": sum(1 for f in findings if f.control_id in mapped_ids),
        })
    tables["framework_coverage"] = pd.DataFrame(frows)

    # initiative prioritisation / risk-vs-effort matrix
    inits = re.prioritize(findings)
    tables["risk_effort_matrix"] = pd.DataFrame([{
        "initiative": r["initiative"],
        "implementation_effort": r["implementation_effort"],
        "effort_rank": r["effort_rank"],
        "expected_risk_reduction_pct": r["avg_risk_reduction_pct"],
        "risk_points_removed": r["risk_points_removed"],
        "total_risk_score": r["total_risk_score"],
        "findings": r["findings"],
        "critical_high": r["critical_high"],
        "quadrant": r["quadrant"],
        "priority_score": r["priority_score"],
        "rank": r["rank"],
    } for r in inits])
    tables["investment_priorities"] = tables["risk_effort_matrix"].head(10).copy()

    # KPIs (key/value)
    kpi = [
        ("genai_readiness_score", readiness["overall"]),
        ("readiness_band", readiness["band"]),
        ("controls_tested", len(config.CONTROLS)),
        ("controls_not_effective",
         sum(1 for r in results if r.status == "Not Effective")),
        ("controls_partially_effective",
         sum(1 for r in results if r.status == "Partially Effective")),
        ("controls_effective", sum(1 for r in results if r.status == "Effective")),
        ("total_findings", len(findings)),
        ("critical_findings", sev.get("Critical", 0)),
        ("high_findings", sev.get("High", 0)),
        ("medium_findings", sev.get("Medium", 0)),
        ("low_findings", sev.get("Low", 0)),
        ("assets_assessed", config.TOTAL_ASSETS_TARGET),
        ("identities_assessed", config.DATASET_SIZES["identities"]),
        ("cloud_resources_assessed", config.DATASET_SIZES["cloud_resources"]),
        ("applications_assessed", config.DATASET_SIZES["applications"]),
        ("data_assets_assessed", config.DATASET_SIZES["data_assets"]),
        ("vendors_assessed", config.DATASET_SIZES["vendors"]),
    ]
    tables["kpis"] = pd.DataFrame(kpi, columns=["metric", "value"])

    return tables, inits


def build_roadmap(findings, inits) -> dict[str, pd.DataFrame]:
    """One initiative per row, bucketed into the earliest phase of its controls."""
    phase_by_initiative: dict[str, str] = {}
    for c in config.CONTROLS:
        cur = phase_by_initiative.get(c["initiative"])
        if cur is None or PHASE_ORDER[c["roadmap_phase"]] < PHASE_ORDER[cur]:
            phase_by_initiative[c["initiative"]] = c["roadmap_phase"]

    rows = []
    for r in inits:
        phase = phase_by_initiative[r["initiative"]]
        rows.append({
            "phase": PHASE_LABEL[phase],
            "phase_order": PHASE_ORDER[phase],
            "priority_rank": r["rank"],
            "initiative": r["initiative"],
            "objective": _initiative_objective(r["initiative"]),
            "domains": r["domains"],
            "linked_findings": r["findings"],
            "critical_high_findings": r["critical_high"],
            "implementation_effort": r["implementation_effort"],
            "expected_risk_reduction_pct": r["avg_risk_reduction_pct"],
            "risk_points_removed": r["risk_points_removed"],
            "owner": r["owner"],
            "quadrant": r["quadrant"],
        })
    df = pd.DataFrame(rows).sort_values(
        ["phase_order", "priority_rank"]).reset_index(drop=True)

    sheets = {"Roadmap (all)": df.drop(columns=["phase_order"])}
    for phase in ["Stabilize", "Govern", "Scale"]:
        sub = df[df["phase"] == PHASE_LABEL[phase]].drop(columns=["phase_order"])
        sheets[phase] = sub.reset_index(drop=True)
    return sheets


def _initiative_objective(name: str) -> str:
    return {
        "Privileged MFA Enforcement":
            "Enforce phishing-resistant MFA on every privileged account; block legacy auth.",
        "Joiner-Mover-Leaver Automation":
            "Automate de-provisioning from HR; weekly reconciliation of leavers vs. active accounts.",
        "Automated Access Reviews":
            "Stand up quarterly access certification and an automated 90-day dormancy policy.",
        "Privileged Access Modernisation":
            "Just-in-time elevation, separate admin accounts and re-scoped service identities.",
        "IAM Governance Uplift":
            "Assign accountable owners to every privileged account and enforce at creation.",
        "Cloud Exposure Remediation":
            "Remove public storage exposure, enable encryption and harden internet-facing services.",
        "Continuous Cloud Control Monitoring":
            "Azure Policy guardrails with drift detection and lifecycle inventory.",
        "Cloud Resilience Uplift":
            "Backup with tested restore for all production data services.",
        "Cloud Network Segmentation":
            "Private endpoints and default-deny network rules across the tenant.",
        "Sensitive-Data Classification":
            "Re-classify and assign owners for all PII-bearing and sensitive datasets.",
        "Sensitive-Data Protection":
            "Encrypt all Confidential/Restricted data with customer-managed keys.",
        "AI Data-Loss Prevention":
            "Prompt-side DLP, retrieval allow-lists and data-minimisation reviews for AI pathways.",
        "Data Retention Governance":
            "Define retention for all datasets and auto-purge AI prompt / conversation data.",
        "AI Governance Operating Model":
            "AI inventory, ownership, risk tiering, approved use cases and human-oversight rules.",
        "AI Audit Logging":
            "Structured logging of AI prompts, responses and retrieval to the SIEM.",
        "AI Output Monitoring":
            "Output-quality, bias and drift monitoring feeding the AI governance forum.",
        "Third-Party AI Risk Review":
            "AI-specific vendor risk reviews, data-handling disclosure and supplier oversight.",
        "SOC Operating Model Uplift":
            "Alert ownership, severity taxonomy and 12-month log retention.",
        "Continuous Control Monitoring":
            "Close detection-coverage gaps on high-value systems and track as a KPI.",
        "Incident Response Readiness":
            "Runbooks for critical systems and an AI-incident playbook with a tabletop exercise.",
        "Vulnerability Management Uplift":
            "Enforce a 30-day critical-vulnerability SLA with executive reporting.",
        "Technology Risk Governance":
            "Risk ownership, remediation dates, exception approvals and contract review cadence.",
    }.get(name, "See linked findings and recommendations.")


# ---------------------------------------------------------------------------
# Excel writers
# ---------------------------------------------------------------------------

def _write_xlsx(path: str, sheets: dict[str, pd.DataFrame],
                severity_col: str | None = None) -> None:
    from openpyxl.styles import PatternFill, Font
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
            ws = writer.sheets[name[:31]]
            for i, col in enumerate(df.columns, 1):
                width = min(60, max(12, int(df[col].astype(str).map(len).max()) + 2))
                ws.column_dimensions[get_column_letter(i)].width = width
            for cell in ws[1]:
                cell.font = Font(bold=True)
            if severity_col and severity_col in df.columns:
                sev_ix = list(df.columns).index(severity_col) + 1
                for row in range(2, len(df) + 2):
                    val = ws.cell(row=row, column=sev_ix).value
                    if val in SEV_FILL:
                        fill = PatternFill("solid", fgColor=SEV_FILL[val])
                        ws.cell(row=row, column=sev_ix).fill = fill


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

def run() -> dict:
    print("[1/5] generating synthetic environment ...")
    dfs = generate_data.main()

    print("[2/5] running 48 controls ...")
    results = control_tests.run_all_controls(dfs)

    print("[3/5] scoring 137 findings ...")
    findings = re.score_findings(control_tests.all_findings(results))
    readiness = re.compute_readiness(findings)

    print("[4/5] building deliverables ...")
    os.makedirs(OUT_DIR, exist_ok=True)
    os.makedirs(PBI_DIR, exist_ok=True)

    register = build_risk_register(findings)
    assessment = build_assessment_results(results)
    crosswalk = build_framework_crosswalk(results)
    pbi_tables, inits = build_powerbi_tables(findings, results, readiness)
    roadmap_sheets = build_roadmap(findings, inits)

    assessment.to_csv(os.path.join(OUT_DIR, "assessment_results.csv"), index=False)
    _write_xlsx(os.path.join(OUT_DIR, "risk_register.xlsx"),
                {"Risk Register": register}, severity_col="severity")
    _write_xlsx(os.path.join(OUT_DIR, "framework_crosswalk.xlsx"),
                {"Control Crosswalk": crosswalk,
                 "Framework Coverage": pbi_tables["framework_coverage"]})
    _write_xlsx(os.path.join(OUT_DIR, "remediation_roadmap.xlsx"), roadmap_sheets)

    exec_summary = _build_exec_summary(findings, results, readiness)
    exec_summary.to_csv(os.path.join(OUT_DIR, "executive_summary.csv"), index=False)

    for name, table in pbi_tables.items():
        table.to_csv(os.path.join(PBI_DIR, f"{name}.csv"), index=False)

    print("[5/5] validating metrics ...")
    _validate(dfs, results, findings, readiness)

    _print_summary(findings, results, readiness, inits)
    return {"findings": findings, "results": results, "readiness": readiness}


def _build_exec_summary(findings, results, readiness) -> pd.DataFrame:
    sev = Counter(f.severity for f in findings)
    rows = [
        ("GenAI Readiness Score", f"{readiness['overall']} / 100"),
        ("Readiness Band", readiness["band"]),
        ("Launch Recommendation",
         "Controlled launch after resolving Critical and High findings"),
        ("Controls Tested", len(config.CONTROLS)),
        ("Controls Not Effective",
         sum(1 for r in results if r.status == "Not Effective")),
        ("Controls Partially Effective",
         sum(1 for r in results if r.status == "Partially Effective")),
        ("Controls Effective", sum(1 for r in results if r.status == "Effective")),
        ("Total Findings", len(findings)),
        ("Critical Findings", sev.get("Critical", 0)),
        ("High Findings", sev.get("High", 0)),
        ("Medium Findings", sev.get("Medium", 0)),
        ("Low Findings", sev.get("Low", 0)),
        ("Assets / Identities Assessed", config.TOTAL_ASSETS_TARGET),
    ]
    for d in config.DOMAINS:
        rows.append((f"Readiness - {d}", f"{readiness['by_domain'][d]} / 100"))
    return pd.DataFrame(rows, columns=["metric", "value"])


def _validate(dfs, results, findings, readiness) -> None:
    total_assets = sum(len(dfs[k]) for k in config.DATASET_SIZES)
    sev = Counter(f.severity for f in findings)

    assert total_assets == 2000, f"total assets {total_assets} != 2000"
    assert len(config.CONTROLS) == 48, "controls != 48"
    assert len(results) == 48, "control results != 48"
    assert len(findings) == 137, f"findings {len(findings)} != 137"
    assert dict(sev) == config.TARGET_SEVERITY_MIX, (
        f"severity mix {dict(sev)} != {config.TARGET_SEVERITY_MIX}")
    lo, hi = config.READINESS_TARGET_BAND
    assert lo <= readiness["overall"] <= hi, (
        f"readiness {readiness['overall']} outside {config.READINESS_TARGET_BAND}")

    expected_files = [
        "assessment_results.csv", "risk_register.xlsx", "framework_crosswalk.xlsx",
        "remediation_roadmap.xlsx", "executive_summary.csv",
    ]
    for fn in expected_files:
        p = os.path.join(OUT_DIR, fn)
        assert os.path.exists(p) and os.path.getsize(p) > 0, f"missing output: {fn}"
    for fn in ["findings", "domain_summary", "severity_summary", "top10_risks",
               "control_status", "framework_coverage", "risk_effort_matrix",
               "investment_priorities", "kpis"]:
        p = os.path.join(PBI_DIR, f"{fn}.csv")
        assert os.path.exists(p) and os.path.getsize(p) > 0, f"missing pbi table: {fn}"


def _print_summary(findings, results, readiness, inits) -> None:
    sev = Counter(f.severity for f in findings)
    status = Counter(r.status for r in results)
    print("\n" + "=" * 66)
    print("  ALL METRICS VALIDATED")
    print("=" * 66)
    print(f"  Assets / identities assessed : {config.TOTAL_ASSETS_TARGET}")
    print(f"  Controls tested              : {len(results)}")
    print(f"    Not Effective              : {status['Not Effective']}")
    print(f"    Partially Effective        : {status['Partially Effective']}")
    print(f"    Effective                  : {status['Effective']}")
    print(f"  Findings                     : {len(findings)} "
          f"(C{sev['Critical']} / H{sev['High']} / M{sev['Medium']} / L{sev['Low']})")
    print(f"  GenAI readiness score        : {readiness['overall']} / 100  "
          f"({readiness['band']})")
    print("\n  Per-domain readiness:")
    for d, s in readiness["by_domain"].items():
        print(f"    {s:5.1f}  {d}")
    print("\n  Top 5 investment priorities:")
    for r in inits[:5]:
        print(f"    {r['rank']}. {r['initiative']} "
              f"[{r['quadrant']}, {r['implementation_effort']} effort]")
    print("\n  Outputs written to outputs/ and outputs/powerbi/")
    print("=" * 66)


if __name__ == "__main__":
    run()
